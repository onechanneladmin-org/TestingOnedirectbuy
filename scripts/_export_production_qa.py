from openpyxl import load_workbook
import json, sys
wb = load_workbook(sys.argv[1], data_only=True)
ws = wb["Use Cases"]
rows = []
for r in ws.iter_rows(min_row=2, values_only=True):
    if not r or not r[0]:
        continue
    rows.append({
        "id": str(r[0]).strip(),
        "section": r[1] or "",
        "subsection": r[2] or "",
        "priority": r[3] or "",
        "title": r[4] or "",
        "surfaces": r[5] or "",
        "desc": r[6] or "",
    })
print(json.dumps(rows))
